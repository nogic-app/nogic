"""Reusable python-docx formatting components for chat exports."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from markdown_clean import markdown_runs
from memo_export_assets import user_initials
from memo_content import chat_message_items, file_items

_ARABIC = re.compile(r"[\u0600-\u06ff\u0750-\u077f\u08a0-\u08ff\ufb50-\ufdff\ufe70-\ufeff]")
ACCENT_SUBTLE = "E8EDE7"
USER_BUBBLE = "F0F0ED"
USER_AVATAR = "F2F1EC"


def set_default_font(document: Any) -> None:
    style = document.styles["Normal"]
    style.font.name = "Segoe UI"
    style.font.size = Pt(9.5)
    properties = style._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    for attribute in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attribute}"), "Segoe UI")


def set_rtl(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    properties = paragraph._p.get_or_add_pPr()
    if properties.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        properties.append(bidi)


def align_for_text(paragraph: Any, text: str) -> None:
    if _ARABIC.search(text):
        set_rtl(paragraph)


def docx_run(
    paragraph: Any,
    text: str,
    size: float = 9.5,
    color: RGBColor = RGBColor(0x6B, 0x6B, 0x66),
    bold: bool = False,
) -> Any:
    run = paragraph.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    fonts.set(qn("w:eastAsia"), "Segoe UI")
    return run


def docx_markdown(
    paragraph: Any,
    text: str,
    size: float,
    color: RGBColor,
) -> None:
    """Write Markdown as real bold/italic runs and preserve list lines."""
    align_for_text(paragraph, text)
    for chunk, bold, italic in markdown_runs(text):
        if chunk == "\n":
            paragraph.add_run().add_break()
            continue
        run = docx_run(paragraph, chunk, size, color, bold)
        run.italic = italic


def shade_cell(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell_margins(cell: Any) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", 100), ("start", 140), ("bottom", 100), ("end", 140)):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    properties.append(margins)


def _remove_table_borders(table: Any) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    properties.append(borders)


def add_conversation(
    document: Any,
    messages: list[dict[str, Any]],
    body_color: RGBColor,
    user_name: str,
    icon_path: Path,
) -> None:
    for message in chat_message_items(messages):
        is_user = message["role"] == "user"
        table = document.add_table(rows=1, cols=2)
        table.alignment = (
            WD_TABLE_ALIGNMENT.RIGHT if is_user else WD_TABLE_ALIGNMENT.LEFT
        )
        table.autofit = False
        table.columns[0].width = Inches(0.42)
        table.columns[1].width = Inches(5.55)
        avatar_cell, bubble_cell = table.rows[0].cells
        avatar_cell.width = Inches(0.42)
        bubble_cell.width = Inches(5.55)
        _remove_table_borders(table)
        avatar_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        avatar = avatar_cell.paragraphs[0]
        avatar.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if is_user:
            shade_cell(avatar_cell, USER_AVATAR)
            docx_run(avatar, user_initials(user_name), 7, body_color, True)
        else:
            avatar.add_run().add_picture(str(icon_path), width=Mm(7), height=Mm(7))

        bubble_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        shade_cell(bubble_cell, USER_BUBBLE if is_user else ACCENT_SUBTLE)
        _set_cell_margins(bubble_cell)
        meta = bubble_cell.paragraphs[0]
        meta.paragraph_format.space_after = Pt(2)
        if is_user:
            meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        speaker = "You" if is_user else "Nogic"
        timestamp = message["timestamp"]
        header = f"{speaker} · {timestamp}" if timestamp else speaker
        docx_run(meta, header, 8, body_color, True)
        body = bubble_cell.add_paragraph()
        body.paragraph_format.space_after = Pt(0)
        docx_markdown(body, message["content"], 9.5, body_color)
        document.add_paragraph().paragraph_format.space_after = Pt(3)


def add_files(
    document: Any,
    messages: list[dict[str, Any]],
    accent: RGBColor,
    body_color: RGBColor,
) -> None:
    files = file_items(messages)
    if not files:
        return
    document.add_page_break()
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(14)
    docx_run(heading, "Files shared in this chat", 12, accent, True)
    for name, file_type, size in files:
        paragraph = document.add_paragraph()
        align_for_text(paragraph, name)
        docx_run(paragraph, f"{name} · {file_type} · {size}", 9.5, body_color)


def add_bottom_border(paragraph: Any, color: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    properties.append(borders)
