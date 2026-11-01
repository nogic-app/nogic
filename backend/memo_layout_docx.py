"""python-docx renderer for chat exports and structured memo sections."""

from __future__ import annotations

from datetime import date
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from memo_docx_components import (
    add_bottom_border,
    add_conversation,
    add_files,
    align_for_text,
    docx_run,
    set_default_font,
    shade_cell,
)
from memo_export_assets import app_icon_path
from memo_content import (
    EMPTY_SECTION,
    chat_message_items,
    memo_sections,
    plant_table_items,
    source_items,
)

INK = RGBColor(0x1F, 0x1F, 0x1D)
BODY = RGBColor(0x2B, 0x2B, 0x28)
SECONDARY = RGBColor(0x6B, 0x6B, 0x66)
ACCENT = RGBColor(0x4A, 0x6B, 0x4E)
RULE = "E5E3DD"


def _add_header(
    document: Any,
    title: str,
    model_name: str,
    message_count: int,
    current_date: str,
) -> None:
    header = document.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    header.autofit = False
    header.columns[0].width = Inches(3.2)
    header.columns[1].width = Inches(3.2)
    left, right = header.rows[0].cells
    left.width = Inches(3.2)
    right.width = Inches(3.2)
    docx_run(left.paragraphs[0], "Nogic", 16, ACCENT, True)
    docx_run(right.paragraphs[0], "Chat export", 9.5, SECONDARY)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)
    align_for_text(heading, title)
    docx_run(heading, title or "Chat", 25, INK, True)

    metadata = document.add_paragraph()
    metadata.paragraph_format.space_after = Pt(12)
    text = (
        f"Exported {current_date} · {message_count} messages · Model {model_name}"
    )
    align_for_text(metadata, text)
    docx_run(metadata, text, 9.5, SECONDARY)
    add_bottom_border(metadata, RULE)


def _add_plant_table(document: Any, messages: list[dict[str, Any]]) -> None:
    plants = plant_table_items(messages)
    if not plants:
        docx_run(document.add_paragraph(), EMPTY_SECTION)
        return
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, label in zip(
        table.rows[0].cells,
        ("Latin name", "Common name", "Water", "Sun", "Source"),
    ):
        shade_cell(cell, "E8EDE7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        docx_run(cell.paragraphs[0], label, 8, ACCENT, True)
    for row in plants:
        for cell, value in zip(table.add_row().cells, row):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            align_for_text(cell.paragraphs[0], value)
            docx_run(cell.paragraphs[0], value, 8, BODY)


def _append_field(paragraph: Any, field: str) -> None:
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    paragraph._p.append(begin_run)
    instruction_run = OxmlElement("w:r")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {field} "
    instruction_run.append(instruction)
    paragraph._p.append(instruction_run)
    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    paragraph._p.append(end_run)


def _add_footer(section: Any, current_date: str) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        int(section.page_width - section.left_margin - section.right_margin),
        WD_TAB_ALIGNMENT.RIGHT,
    )
    docx_run(paragraph, f"Exported from Nogic · {current_date}", 8, SECONDARY)
    docx_run(paragraph, "\t", 8, SECONDARY)
    docx_run(paragraph, "Page ", 8, SECONDARY)
    _append_field(paragraph, "PAGE")
    docx_run(paragraph, " of ", 8, SECONDARY)
    _append_field(paragraph, "NUMPAGES")


def _update_fields_on_open(document: Any) -> None:
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _add_sources(document: Any, messages: list[dict[str, Any]]) -> None:
    document.add_page_break()
    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(4)
    docx_run(heading, "Sources", 24, INK, True)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    docx_run(
        subtitle,
        "Every source cited in this chat, in order of first use.",
        9.5,
        SECONDARY,
    )
    sources = source_items(messages)
    if not sources:
        docx_run(document.add_paragraph(), EMPTY_SECTION)
        return
    for index, (label, details) in enumerate(sources, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        align_for_text(paragraph, label)
        docx_run(paragraph, f"[{index}] {label}", 10.5, INK, True)
        detail = details[0] if details else "Source details unavailable."
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.space_after = Pt(12)
        align_for_text(paragraph, detail)
        docx_run(paragraph, detail, 9.5, BODY)


def _add_memo_sections(document: Any, messages: list[dict[str, Any]]) -> None:
    for heading_text, items in memo_sections(messages):
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(16)
        docx_run(heading, heading_text, 12, ACCENT, True)
        if heading_text == "Plant Palette":
            _add_plant_table(document, messages)
        elif not items:
            docx_run(document.add_paragraph(), EMPTY_SECTION)
        else:
            for lines in items:
                lead = document.add_paragraph()
                lead.paragraph_format.space_after = Pt(1)
                docx_run(lead, lines[0], 10.5, INK, True)
                for line in lines[1:]:
                    detail = document.add_paragraph()
                    detail.paragraph_format.left_indent = Pt(18)
                    docx_run(detail, line, 9.5, BODY)


def render_docx(
    title: str,
    messages: list[dict[str, Any]],
    path: str,
    model_name: str,
    user_name: str = "",
) -> None:
    """Render the chat and its memo references to a Word document."""
    current_date = date.today().strftime("%B %d, %Y").replace(" 0", " ")
    document = Document()
    set_default_font(document)
    document.core_properties.title = f"Nogic Chat Export - {title}"
    _update_fields_on_open(document)
    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(22)
        section.right_margin = Mm(22)
        section.top_margin = Mm(17)
        section.bottom_margin = Mm(20)
    _add_header(
        document, title, model_name, len(chat_message_items(messages)), current_date
    )
    add_conversation(document, messages, BODY, user_name, app_icon_path())
    add_files(document, messages, ACCENT, BODY)
    _add_memo_sections(document, messages)
    _add_sources(document, messages)
    for section in document.sections:
        _add_footer(section, current_date)
    document.save(path)
