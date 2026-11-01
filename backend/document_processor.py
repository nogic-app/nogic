"""Local document and image preparation for attachment Q&A."""

import base64
import io
from pathlib import Path
from typing import Any

from PIL import Image

MAX_IMAGE_PIXELS = 256
MAX_ATTACHMENT_BYTES = 25 * 1024 * 1024
TEXT_ATTACHMENT_TYPES = {".pdf", ".docx", ".xlsx"}
IMAGE_ATTACHMENT_TYPES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}


def chunk_text(text: str, max_chars: int = 6000) -> list[str]:
    """Split text near sentence or line boundaries without dropping content."""
    if max_chars <= 0:
        raise ValueError("max_chars must be positive.")
    remaining = text.strip()
    chunks: list[str] = []
    while remaining:
        if len(remaining) <= max_chars:
            chunks.append(remaining)
            break
        boundary = max(
            remaining.rfind("\n", 0, max_chars + 1),
            remaining.rfind(". ", 0, max_chars + 1),
            remaining.rfind(".\n", 0, max_chars + 1),
        )
        split_at = boundary + 1 if boundary >= max_chars // 2 else max_chars
        chunks.append(remaining[:split_at].strip())
        remaining = remaining[split_at:].strip()
    return chunks


def _extract_pdf(path: Path) -> str:
    import pymupdf

    with pymupdf.open(path) as document:
        return "\n".join(page.get_text() for page in document)


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return "\n".join(parts)


def _excel_markdown(path: Path) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets: list[str] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            width = max(len(row) for row in rows)
            normalized = [
                ["" if value is None else str(value) for value in row] + [""] * (width - len(row))
                for row in rows
            ]
            header = normalized[0]
            lines = [
                f"## Sheet: {sheet.title}",
                "| " + " | ".join(header) + " |",
                "| " + " | ".join("---" for _ in header) + " |",
            ]
            lines.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
            sheets.append("\n".join(lines))
        return sheets
    finally:
        workbook.close()


def extract_document_chunks(path: Path, max_chars: int = 6000) -> list[str]:
    """Extract PDF/Word/Excel content and return bounded Q&A chunks."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return chunk_text(_extract_pdf(path), max_chars)
    if suffix == ".docx":
        return chunk_text(_extract_docx(path), max_chars)
    if suffix == ".xlsx":
        chunks: list[str] = []
        for sheet in _excel_markdown(path):
            chunks.extend(chunk_text(sheet, max_chars))
        return chunks
    return []


def prepare_image(path: Path) -> dict[str, Any]:
    """Resize an image to 256px maximum and encode it as a compact JPEG."""
    with Image.open(path) as image:
        image.thumbnail((MAX_IMAGE_PIXELS, MAX_IMAGE_PIXELS), Image.Resampling.LANCZOS)
        if image.mode not in {"RGB", "L"}:
            image = image.convert("RGB")
        output = io.BytesIO()
        image.save(output, format="JPEG", quality=80, optimize=True)
    return {
        "name": path.name,
        "data": base64.b64encode(output.getvalue()).decode("ascii"),
    }


def process_attachments(
    raw_attachments: object,
    destination: Path,
) -> tuple[list[dict[str, object]], list[str], list[dict[str, Any]]]:
    """Copy attachments and return metadata, document chunks, and image data."""
    if raw_attachments is None:
        return [], [], []
    if not isinstance(raw_attachments, list):
        raise ValueError("Attachments must be a list.")

    saved: list[dict[str, object]] = []
    chunks: list[str] = []
    images: list[dict[str, Any]] = []
    for raw_attachment in raw_attachments:
        if not isinstance(raw_attachment, dict) or not isinstance(raw_attachment.get("path"), str):
            raise ValueError("Each attachment must include a file path.")
        source = Path(raw_attachment["path"]).expanduser()
        if not source.is_file():
            raise ValueError(f"File not found: {source}")
        if source.stat().st_size > MAX_ATTACHMENT_BYTES:
            raise ValueError(f"Attachment exceeds the 25 MB limit: {source.name}")
        target = destination / source.name
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(source.read_bytes())
        saved.append({
            "name": target.name,
            "path": str(target),
            "type": target.suffix.removeprefix(".").upper() or "FILE",
            "size": target.stat().st_size,
        })
        if target.suffix.lower() in TEXT_ATTACHMENT_TYPES:
            chunks.extend(extract_document_chunks(target))
        elif target.suffix.lower() in IMAGE_ATTACHMENT_TYPES:
            images.append(prepare_image(target))
    return saved, chunks, images
